"""Excel and PDF export for any report produced by registry.run()."""
import io
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from accounts.models import BusinessSettings
from common.money import indian_format

INK = "12344D"


def _cell_value(value, kind):
    if value is None or value == "":
        return ""
    if kind == "money":
        return float(value)
    if kind == "number" and isinstance(value, (int, float, Decimal)):
        return float(value)
    return value


def to_excel(report):
    biz = BusinessSettings.load()
    wb = Workbook()
    ws = wb.active
    ws.title = report["title"][:30]
    cols = report["columns"]
    ws.append([biz.name])
    ws["A1"].font = Font(bold=True, size=14, color=INK)
    ws.append([report["title"]])
    ws["A2"].font = Font(bold=True, size=12)
    ws.append([" | ".join([report["period"]] + report.get("filters_applied", []))])
    row = 5
    for item in report.get("summary", []):
        ws.cell(row=row, column=1, value=item["label"])
        c = ws.cell(row=row, column=2, value=float(item["value"]))
        c.number_format = "#,##0.00"
        row += 1
    if report.get("summary"):
        row += 1

    header_fill = PatternFill("solid", fgColor=INK)
    thin = Side(style="thin", color="D0D5DD")
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=i, value=c["label"])
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="right" if c["type"] in ("money", "number") else "left")
    header_row = row
    for r in report["rows"]:
        row += 1
        style = r.get("_style")
        for i, c in enumerate(cols, start=1):
            cell = ws.cell(row=row, column=i, value=_cell_value(r.get(c["key"]), c["type"]))
            if c["type"] == "money":
                cell.number_format = "#,##0.00"
            elif c["type"] == "date" and isinstance(cell.value, date):
                cell.number_format = "DD-MM-YYYY"
            if style in ("section", "subtotal"):
                cell.font = Font(bold=True)
            cell.border = Border(bottom=thin)
    totals = report.get("totals") or {}
    if totals:
        row += 1
        for i, c in enumerate(cols, start=1):
            v = totals.get(c["key"])
            if i == 1 and v is None:
                v = "Total"
            cell = ws.cell(row=row, column=i, value=_cell_value(v, c["type"]) if v is not None else "")
            cell.font = Font(bold=True)
            cell.border = Border(top=Side(style="medium", color=INK))
            if c["type"] == "money":
                cell.number_format = "#,##0.00"
    if report.get("note"):
        ws.cell(row=row + 2, column=1, value=report["note"]).font = Font(italic=True, color="667085")

    for i, c in enumerate(cols, start=1):
        width = max([len(str(c["label"]))] + [len(str(r.get(c["key"], "") or "")) for r in report["rows"][:500]]) + 3
        ws.column_dimensions[get_column_letter(i)].width = min(max(width, 10), 45)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _fmt(value, kind):
    if value is None or value == "":
        return ""
    if kind == "money":
        return indian_format(value, symbol="")
    if kind == "date" and isinstance(value, date):
        return value.strftime("%d-%m-%Y")
    if kind == "number" and isinstance(value, Decimal):
        return f"{value:,.2f}".rstrip("0").rstrip(".")
    return str(value)


def to_pdf(report):
    biz = BusinessSettings.load()
    buf = io.BytesIO()
    cols = report["columns"]
    wide = len(cols) > 5
    page = landscape(A4) if wide else A4
    doc = SimpleDocTemplate(buf, pagesize=page, leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=report["title"])
    ink = colors.HexColor("#" + INK)
    body = ParagraphStyle("body", fontName="Helvetica", fontSize=8, leading=10)
    story = [
        Paragraph(biz.name, ParagraphStyle("biz", fontName="Helvetica-Bold", fontSize=14, textColor=ink, leading=18)),
    ]
    if biz.address or biz.phone or biz.gst_number:
        meta = " | ".join(x for x in [biz.address.replace("\n", ", "), biz.phone,
                                      f"GST: {biz.gst_number}" if biz.gst_number else ""] if x)
        story.append(Paragraph(meta, ParagraphStyle("meta", fontName="Helvetica", fontSize=8, textColor=colors.grey)))
    story += [
        Spacer(1, 4 * mm),
        Paragraph(report["title"], ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=12, leading=15)),
        Paragraph(" | ".join([report["period"]] + report.get("filters_applied", [])) + " | Amounts in Rs.",
                  ParagraphStyle("p", fontName="Helvetica", fontSize=9, textColor=colors.grey, leading=12)),
        Spacer(1, 4 * mm),
    ]
    if report.get("summary"):
        data = [[i["label"] for i in report["summary"]], [indian_format(i["value"], "") for i in report["summary"]]]
        st = Table(data, hAlign="LEFT")
        st.setStyle(TableStyle([("FONT", (0, 0), (-1, 0), "Helvetica", 8), ("TEXTCOLOR", (0, 0), (-1, 0), colors.grey),
                                ("FONT", (0, 1), (-1, 1), "Helvetica-Bold", 10), ("RIGHTPADDING", (0, 0), (-1, -1), 14)]))
        story += [st, Spacer(1, 4 * mm)]

    header = [Paragraph(f"<b>{c['label']}</b>", ParagraphStyle(
        "h", parent=body, textColor=colors.white, alignment=2 if c["type"] in ("money", "number") else 0)) for c in cols]
    data = [header]
    styles = [
        ("BACKGROUND", (0, 0), (-1, 0), ink),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONT", (0, 1), (-1, -1), "Helvetica", 8),
        ("LINEBELOW", (0, 1), (-1, -1), 0.25, colors.HexColor("#D0D5DD")),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for i, c in enumerate(cols):
        if c["type"] in ("money", "number"):
            styles.append(("ALIGN", (i, 1), (i, -1), "RIGHT"))
    for r in report["rows"]:
        data.append([Paragraph(_fmt(r.get(c["key"]), c["type"]).replace("  ", "&nbsp;&nbsp;"), body)
                     if c["type"] == "text" else _fmt(r.get(c["key"]), c["type"]) for c in cols])
        if r.get("_style") in ("section", "subtotal"):
            styles.append(("FONT", (0, len(data) - 1), (-1, len(data) - 1), "Helvetica-Bold", 8))
    totals = report.get("totals") or {}
    if totals:
        row = []
        for i, c in enumerate(cols):
            v = totals.get(c["key"])
            row.append("Total" if i == 0 and v is None else _fmt(v, c["type"]))
        data.append(row)
        styles += [("FONT", (0, len(data) - 1), (-1, len(data) - 1), "Helvetica-Bold", 8),
                   ("LINEABOVE", (0, len(data) - 1), (-1, len(data) - 1), 1, ink)]
    if len(data) == 1:
        data.append(["No records for this period."] + [""] * (len(cols) - 1))
    table = Table(data, repeatRows=1, hAlign="LEFT", colWidths=_widths(cols, page[0] - 24 * mm))
    table.setStyle(TableStyle(styles))
    story.append(table)
    if report.get("note"):
        story += [Spacer(1, 3 * mm), Paragraph(report["note"], ParagraphStyle("n", parent=body, textColor=colors.grey))]

    def footer(canvas, d):
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.grey)
        canvas.drawString(12 * mm, 7 * mm, f"Generated on {report['generated']:%d %b %Y}")
        canvas.drawRightString(page[0] - 12 * mm, 7 * mm, f"Page {d.page}")

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def _widths(cols, total):
    weights = [{"money": 1.1, "number": 0.8, "date": 0.9, "status": 1.0}.get(c["type"], 1.6) for c in cols]
    s = sum(weights)
    return [total * w / s for w in weights]
